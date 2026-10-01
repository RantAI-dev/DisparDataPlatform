# -*- coding: utf-8 -*-
"""
Generator laporan bulan ke-4 September 2026.
Periode pekerjaan: 1–30 September. Verifikasi: 2 Oktober 2026 WIB.
Format dan identitas tetap mengikuti laporan sebelumnya.
Angka dan SQL sumber disimpan dalam verifikasi-bulan4.json.
"""
import os
import json
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image

BASE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(BASE, "assets-bulan4")
INK = RGBColor(0x1a, 0x1a, 0x1a)
GREY = RGBColor(0x55, 0x55, 0x55)

# ---------- helpers ----------

def set_cell_bg(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def set_table_borders(table, color="7f7f7f", sz="4"):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{edge}')
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), sz)
        e.set(qn('w:space'), '0'); e.set(qn('w:color'), color)
        borders.append(e)
    tblPr.append(borders)

def cell_text(cell, text, bold=False, size=10, color=INK, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align: p.alignment = align
    run = p.add_run(text)
    run.bold = bold; run.font.size = Pt(size); run.font.name = "Arial"; run.font.color.rgb = color
    return p

def para(doc, text, size=11, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
         color=INK, space_after=6, space_before=0):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after); pf.space_before = Pt(space_before); pf.line_spacing = 1.15
    run = p.add_run(text)
    run.bold = bold; run.italic = italic; run.font.size = Pt(size)
    run.font.name = "Arial"; run.font.color.rgb = color
    return p

def heading(doc, text, size=12):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12); p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True; run.font.size = Pt(size); run.font.name = "Arial"; run.font.color.rgb = INK
    return p

def bullets(doc, items, letters=False):
    for i, it in enumerate(items):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.9); p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.12
        prefix = f"{chr(97+i)}.  " if letters else "•  "
        r = p.add_run(prefix); r.font.name = "Arial"; r.font.size = Pt(11); r.bold = letters
        r2 = p.add_run(it); r2.font.name = "Arial"; r2.font.size = Pt(11); r2.font.color.rgb = INK

def add_image(doc, filename, caption, width_cm=15.5, max_h_cm=19.0):
    path = os.path.join(ASSETS, filename)
    w_px, h_px = Image.open(path).size
    w = width_cm
    h = width_cm * h_px / w_px
    if h > max_h_cm:
        h = max_h_cm; w = max_h_cm * w_px / h_px
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(path, width=Cm(w), height=Cm(h))
    cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(10)
    r = cap.add_run(caption); r.italic = True; r.font.size = Pt(9)
    r.font.name = "Arial"; r.font.color.rgb = GREY

def realisasi_table(doc, rows):
    t = doc.add_table(rows=1, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    set_table_borders(t)
    hdr = t.rows[0].cells
    cell_text(hdr[0], "Uraian Tugas (KAK)", bold=True, size=10, align=WD_ALIGN_PARAGRAPH.LEFT)
    cell_text(hdr[1], "Realisasi Bulan Ini", bold=True, size=10, align=WD_ALIGN_PARAGRAPH.LEFT)
    for c in hdr: set_cell_bg(c, "d9d9d9")
    for tugas, real in rows:
        r = t.add_row().cells
        cell_text(r[0], tugas, size=9.5, align=WD_ALIGN_PARAGRAPH.LEFT)
        cell_text(r[1], real, size=9.5, align=WD_ALIGN_PARAGRAPH.LEFT)
    t.columns[0].width = Cm(6.6); t.columns[1].width = Cm(9.0)
    for row in t.rows:
        row.cells[0].width = Cm(6.6); row.cells[1].width = Cm(9.0)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def data_table(doc, header, rows, widths):
    t = doc.add_table(rows=1, cols=len(header)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False; set_table_borders(t)
    for i, h in enumerate(header):
        cell_text(t.rows[0].cells[i], h, bold=True, size=8.5, align=WD_ALIGN_PARAGRAPH.LEFT)
        set_cell_bg(t.rows[0].cells[i], "d9d9d9")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cell_text(cells[i], v, size=8.5, align=WD_ALIGN_PARAGRAPH.LEFT)
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


# ---------- bukti terverifikasi, dengan periode eksplisit ----------
with open(os.path.join(BASE, "verifikasi-bulan4.json"), encoding="utf-8") as f:
    EVIDENCE = json.load(f)
INV = EVIDENCE["inventory"]
QUALITY = EVIDENCE["quality_september_last"]
HISTORY = EVIDENCE["rowcount_history"]
QUARANTINE = EVIDENCE["quarantine"]

def angka(value):
    return f"{value:,}".replace(",", ".")

LAKE_HEADER = ["Lapisan", "Bentuk penyimpanan", "Isi", "Jumlah"]
LAKE_ROWS = [
    ["Data mentah", "Format terbuka Apache Iceberg", "Tabel data primer, berkas, dan sekunder; tidak termasuk metadata", f'{INV["bronze_data_tables"]} tabel data'],
    ["Metadata mentah", "Tabel pendukung", "Metadata sumber; dihitung terpisah dari tabel data", str(INV["bronze_objects_including_metadata"] - INV["bronze_data_tables"]) + " tabel"],
    ["Data bersih", "Tampilan dan tabel bertipe", "Tampilan bersih serta tabel pendukung di silver", f'{INV["silver_views"]} tampilan + {INV["silver_tables"]} tabel'],
    ["Mart penyaji", "Tabel agregat", "Mart operasional; tabel berakhiran _baru tidak dihitung", f'{INV["gold_operational_marts"]} mart'],
    ["Persiapan penyaji", "Tabel berakhiran _baru", "Objek persiapan; dipisahkan dari mart operasional", str(INV["gold_staging_baru"]) + " tabel"],
]
LAKE_WIDTHS = [3.0, 3.5, 6.1, 3.0]
MART_HEADER = ["Mart penyaji", "Isi / batas penggunaan", "Baris"]
MART_DESC = {
    "mart_wisman": "Arsip kunjungan menurut kawasan; seluruh data tahun 2014",
    "mart_kunjungan_dtw": "Kunjungan daya tarik wisata; angka terisi berperiode Juni/Juli 2026, bukan September",
    "mart_gci_readiness": "Kesiapan data indikator GCI/GPCI",
    "mart_event": "Ringkasan event",
    "mart_kuliner": "Ringkasan kuliner",
    "mart_atlas": "Ringkasan pendataan lapangan",
    "mart_hotel_scraped_booking": "Hasil uji pengumpulan data hotel; bukan data statistik resmi",
}
MART_ROWS = [[name, MART_DESC[name], angka(n)] for name, n in EVIDENCE["marts"]["results"].items()]
MART_WIDTHS = [4.7, 8.9, 2.0]
QUAL_HEADER = ["Pemeriksaan / riwayat", "Hasil dan periode"]
QUAL_ROWS = [
    ["Lolos", angka(QUALITY["pass"]) + " pemeriksaan; 30 September 2026"],
    ["Peringatan", angka(QUALITY["warn"]) + " pemeriksaan; 30 September 2026"],
    ["Gagal", angka(QUALITY["fail"]) + " pemeriksaan; 30 September 2026"],
    ["Catatan jumlah baris selama September", angka(HISTORY["september_records"]) + " catatan pada " + str(HISTORY["september_distinct_tables"]) + " tabel"],
    ["Observasi jumlah baris terakhir September", "30 September 2026 09.30.49 WIB; " + str(HISTORY["last_september_run_records"]) + " catatan"],
]
QUAL_WIDTHS = [8.0, 7.6]
KARANTINA_HEADER = ["Alasan / keterangan", "Tabel terdampak", "Temuan"]
KARANTINA_ROWS = [[r["reason"], str(r["tables"]) + " tabel", angka(r["failed_values"]) + " nilai"] for r in QUARANTINE["samples"]]
KARANTINA_ROWS.append(["Total kelompok temuan snapshot 1 Oktober", "Bukan baris/orang unik", str(QUARANTINE["snapshot_groups"]) + " kelompok"])
KARANTINA_WIDTHS = [8.0, 3.6, 4.0]
WIS_HEADER = ["Indikator", "Periode data", "Jumlah"]
WIS_ROWS = [
    ["Wisman Jakarta", "Januari–Desember 2025", angka(EVIDENCE["indicators"][0]["value"]) + " kunjungan"],
    ["Wisman Jakarta", "Januari–April 2026", angka(next(r["jumlah"] for r in EVIDENCE["indicators"][0]["results"] if r["tahun"] == 2026)) + " kunjungan"],
    ["Perjalanan wisnus ke kota tujuan DKI", "Januari–Desember 2025", angka(EVIDENCE["indicators"][1]["values"]["2025"]) + " perjalanan"],
    ["Perjalanan wisnus ke kota tujuan DKI", "Januari–Mei 2026", angka(EVIDENCE["indicators"][1]["values"]["2026_Jan_Mei"]) + " perjalanan"],
    ["Tenaga kerja ekonomi kreatif DKI", "2024; kategori Total", angka(EVIDENCE["indicators"][2]["value"]) + " orang"],
]
WIS_WIDTHS = [5.5, 5.0, 5.1]
TABLES = {
    "lake": (LAKE_HEADER, LAKE_ROWS, LAKE_WIDTHS, "Sumber Q1: SHOW TABLES pada lake, silver, serving; SHOW CREATE TABLE pada silver. Snapshot 2 Oktober 2026, bukan inventaris akhir September. Data mentah menghitung bronze_sdi, bronze_file, bronze_sec; metadata dipisah; mart operasional tidak mencakup _baru."),
    "mart": (MART_HEADER, MART_ROWS, MART_WIDTHS, "Sumber Q2: SELECT count() pada tujuh mart serving, 2 Oktober 2026. Jumlah baris adalah isi tabel saat pemeriksaan; tahun dan cakupan data mengikuti keterangan tiap mart."),
    "qual": (QUAL_HEADER, QUAL_ROWS, QUAL_WIDTHS, "Sumber Q3: _silver_meta.quality, batas 30 September WIB; Q4: _silver_meta.rowcount_history, periode 1–30 September WIB. Dibaca 2 Oktober 2026. Catatan tidak membuktikan semua jadwal berhasil atau semua temuan otomatis dikeluarkan dari mart."),
    "karantina": (KARANTINA_HEADER, KARANTINA_ROWS, KARANTINA_WIDTHS, "Sumber Q5: _silver_meta.karantina, snapshot tersimpan 1 Oktober, dibaca 2 Oktober 2026. Snapshot akhir September tidak tersedia. Nilai gagal konversi dan kelompok temuan bukan jumlah baris sumber atau orang unik."),
    "wis": (WIS_HEADER, WIS_ROWS, WIS_WIDTHS, "Sumber Q6: silver.wisman_jakarta_per_bulan, jumlah_kunjungan per tahun; Q7: silver.wisnus_perjalanan_per_kota_tujuan, jumlah_perjalanan per tahun; Q8: silver.tenaga_kerja_ekraf_per_provinsi, tahun 2024, kode_wilayah 31, jenis_kelamin Total. Dibaca 2 Oktober 2026. Angka bukan realisasi bulan September 2026."),
}
PENDAHULUAN_UMUM = (
    "Laporan bulanan keempat ini memuat pelaksanaan kegiatan Penyediaan dan Pengelolaan Data Statistik "
    "Pariwisata Jakarta Tahun Anggaran 2026 untuk periode 1–30 September. Fokus pekerjaan adalah "
    "pengembangan dashboard statistik pariwisata dan ekonomi kreatif, integrasi sumber resmi, perluasan "
    "inventori sekunder, serta dokumentasi operasional. Capaian ditelusuri dari perubahan yang tercatat "
    "pada September dan bukti sistem yang diperiksa pada 2 Oktober 2026. Tahun data indikator dibedakan "
    "dari periode pekerjaan; snapshot inventaris setelah akhir bulan diberi tanggal sebenarnya."
)
RTL_UMUM = (
    "Rencana Oktober adalah melengkapi data yang masih kosong, menyelaraskan definisi dan periode "
    "indikator, menindaklanjuti temuan mutu, serta menguji kemampuan operasional dan analitis yang belum "
    "memiliki bukti memadai. Pengembangan platform AI-Data diteruskan dengan menjaga keterlacakan "
    "angka sampai sumbernya. Data parsial, inventori usaha, dan hasil uji pengumpulan data tidak "
    "dipakai sebagai pengganti statistik resmi yang belum tersedia."
)
CAPTION_URLS = {
    "01-dashboard.jpg": "https://dispar.rantai.dev/dashboard",
    "02-wisman.jpg": "https://dispar.rantai.dev/dashboard/wisman",
    "03-wisnus.jpg": "https://dispar.rantai.dev/dashboard/wisnus",
    "04-pdrb-ekraf.jpg": "https://dispar.rantai.dev/dashboard/pdrb-ekraf",
    "05-tenaga-kerja.jpg": "https://dispar.rantai.dev/dashboard/tenaga-kerja",
    "06-katalog.jpg": "https://dispar.rantai.dev/sdi",
    "07-wellness.jpg": "https://dispar.rantai.dev/sdi/wellness-jakarta",
    "08-kuliner.jpg": "https://dispar.rantai.dev/gci/pariwisata/kuliner-michelin",
    "09-ai.jpg": "https://dispar.rantai.dev/ai",
    "10-api-docs.jpg": "https://dispar.rantai.dev/docs",
    "11-buku-home.jpg": "https://dispar-buku.vercel.app/",
}

def caption(filename, text):
    return text + " Sumber: " + CAPTION_URLS[filename] + "; tangkapan layar 2 Oktober 2026."


def prepare_layout(doc):
    # Jaga baris tabel dan blok tanda tangan tetap utuh pada hasil cetak.
    tables = doc.tables
    for i, table in enumerate(tables):
        for row in table.rows:
            props = row._tr.get_or_add_trPr()
            props.append(OxmlElement("w:cantSplit"))
        if i < len(tables) - 1 and table.rows:
            table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    signature = doc.tables[-1]
    paragraphs = [p for row in signature.rows for cell in row.cells for p in cell.paragraphs]
    for p in paragraphs[:-1]:
        p.paragraph_format.keep_with_next = True

def build(doc_meta):
    doc = Document()
    st = doc.styles['Normal']; st.font.name = "Arial"; st.font.size = Pt(11)
    for section in doc.sections:
        section.top_margin = Cm(2.2); section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.5); section.right_margin = Cm(2.2)

    # KOP
    para(doc, "PEMERINTAH PROVINSI DAERAH KHUSUS IBUKOTA JAKARTA", size=12, bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    para(doc, "DINAS PARIWISATA DAN EKONOMI KREATIF", size=13, bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    para(doc, "Bidang Data, Informasi dan Pengembangan Destinasi", size=10, italic=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, color=GREY, space_after=2)
    hr = doc.add_paragraph(); hr.paragraph_format.space_after = Pt(10)
    pPr = hr._p.get_or_add_pPr(); pbdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom'); bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '12')
    bottom.set(qn('w:space'), '1'); bottom.set(qn('w:color'), '1a1a1a'); pbdr.append(bottom); pPr.append(pbdr)

    para(doc, "LAPORAN PELAKSANAAN TUGAS TENAGA AHLI", size=14, bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1)
    para(doc, "Kegiatan Penyediaan dan Pengelolaan Data Statistik Pariwisata Jakarta — Tahun Anggaran 2026",
         size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, color=GREY, space_after=10)

    # IDENTITAS
    t = doc.add_table(rows=0, cols=2); set_table_borders(t, color="bfbfbf")
    ident = [
        ("Nama Tenaga Ahli", "[Nama Tenaga Ahli]"),
        ("NIK / No. Kontrak", "[NIK / No. Kontrak]"),
        ("Jabatan / Posisi", doc_meta["jabatan"]),
        ("Kualifikasi", doc_meta["kualifikasi"]),
        ("Periode Laporan", "Bulan ke-4 (Laporan Bulanan Keempat) — September 2026"),
        ("Sub Kegiatan", "Perencanaan Daya Tarik Wisata Provinsi"),
        ("Lokasi", "Provinsi DKI Jakarta"),
    ]
    for k, v in ident:
        row = t.add_row().cells
        cell_text(row[0], k, bold=True, size=10, align=WD_ALIGN_PARAGRAPH.LEFT); set_cell_bg(row[0], "eeeeee")
        cell_text(row[1], v, size=10, align=WD_ALIGN_PARAGRAPH.LEFT)
        row[0].width = Cm(4.8); row[1].width = Cm(10.8)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    heading(doc, "I.  PENDAHULUAN")
    para(doc, PENDAHULUAN_UMUM)
    para(doc, doc_meta["pendahuluan_peran"])

    heading(doc, "II.  URAIAN TUGAS")
    para(doc, "Sesuai Kerangka Acuan Kerja (KAK) kegiatan, uraian tugas untuk posisi ini adalah sebagai berikut:",
         space_after=4)
    bullets(doc, doc_meta["uraian_tugas"], letters=True)

    heading(doc, "III.  PELAKSANAAN PEKERJAAN BULAN INI")
    para(doc, doc_meta["pelaksanaan_intro"])
    realisasi_table(doc, doc_meta["realisasi"])

    heading(doc, "IV.  BUKTI PELAKSANAAN")
    para(doc, doc_meta["bukti_intro"], space_after=6)
    for item in doc_meta["bukti"]:
        kind = item[0]
        if kind == "img":
            add_image(doc, item[1], caption(item[1], item[2]))
        else:
            spec = TABLES[kind]
            para(doc, item[1], bold=True, size=10.5, space_before=4, space_after=3,
                 align=WD_ALIGN_PARAGRAPH.LEFT)
            data_table(doc, spec[0], spec[1], spec[2])
            para(doc, spec[3], size=8.5, italic=True, color=GREY,
                 align=WD_ALIGN_PARAGRAPH.LEFT, space_after=8)

    heading(doc, "V.  RENCANA TINDAK LANJUT — PENGEMBANGAN PLATFORM AI-DATA")
    para(doc, RTL_UMUM)
    para(doc, doc_meta["rtl_peran"])
    for item in doc_meta.get("rtl_bukti", []):
        add_image(doc, item[0], item[1])

    heading(doc, "VI.  PENUTUP")
    para(doc, doc_meta["penutup"])

    # tanda tangan
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    sig = doc.add_table(rows=0, cols=2)
    r = sig.add_row().cells
    cell_text(r[0], "Mengetahui,\nPejabat Pembuat Komitmen", size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    cell_text(r[1], "Jakarta, ......................... 2026\nTenaga Ahli yang bersangkutan",
              size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    for _ in range(4): sig.add_row()
    r = sig.add_row().cells
    cell_text(r[0], "(  Bima Agung  )", bold=True, size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    cell_text(r[1], "(  [Nama Tenaga Ahli]  )", bold=True, size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = sig.add_row().cells
    cell_text(r[0], "NIP. 197907162011011008", size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, color=GREY)
    cell_text(r[1], doc_meta["jabatan"], size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, color=GREY)

    outdir = os.path.join(BASE, "bulan-4")
    os.makedirs(outdir, exist_ok=True)
    out = os.path.join(outdir, doc_meta["filename"] + ".docx")
    prepare_layout(doc)
    doc.save(out)
    print("wrote", out)

MEMBERS = [
{
    "filename": "01-PM-Analisis-Data-Dashboard",
    "jabatan": "Project Manager Analisis Data dan Pembangunan Dashboard Pariwisata",
    "kualifikasi": "S2 / setara, pengalaman minimal 5 tahun",
    "uraian_tugas": [
        "Mengelola dan merencanakan SDM yang dibutuhkan untuk kegiatan Analisis Data dan Pembangunan Dashboard Pariwisata, termasuk analisis, desain, pelaksanaan, pengujian/validasi, dan output yang dihasilkan;",
        "Memberikan laporan kegiatan Analisis Data dan Pembangunan Dashboard Pariwisata berupa progres pekerjaan dan pembaruan;",
        "Memastikan kegiatan selesai tepat waktu, sesuai anggaran, memenuhi standar kualitas, dan harapan pemilik program;",
        "Mengelola risiko dan isu pekerjaan, memastikan penyediaan informasi tepat waktu, serta melakukan mitigasi risiko dan langkah eskalasi."
    ],
    "pendahuluan_peran": "Pada bulan keempat, Project Manager mengarahkan perluasan platform dari katalog data menuju dashboard statistik pariwisata dan ekonomi kreatif. Prioritas pekerjaan mencakup pemanfaatan data resmi, pengayaan data sekunder, dan penyelarasan dokumentasi operasional agar keluaran dapat diperiksa oleh Dinas.",
    "pelaksanaan_intro": "Realisasi September ditelusuri melalui riwayat perubahan repositori, dokumentasi operasional, serta pemeriksaan halaman publik. Pemutakhiran yang terjadi setelah 30 September tidak dihitung sebagai capaian periode ini.",
    "realisasi": [
        [
            "Perencanaan dan pembagian pekerjaan",
            "Mengelompokkan pekerjaan September ke pengembangan dashboard statistik, integrasi sumber resmi, pengayaan inventori sekunder, dan dokumentasi pemanfaatan platform."
        ],
        [
            "Pelaporan progres dan pembaruan",
            "Merangkum pengembangan dashboard wisatawan mancanegara, wisatawan nusantara, PDRB ekonomi kreatif, dan tenaga kerja ekonomi kreatif yang tercatat pada September dan dapat diperiksa pada platform publik."
        ],
        [
            "Pengendalian mutu keluaran",
            "Menetapkan pemisahan antara tahun data indikator, periode pekerjaan, dan tanggal verifikasi. Data parsial dan data yang belum tersedia dinyatakan pada halaman masing-masing."
        ],
        [
            "Pengelolaan risiko pekerjaan",
            "Mencatat kebutuhan data lama menginap non-bintang, PDRB per wilayah, serta rincian tenaga kerja per subsektor DKI sebagai tindak lanjut. Pengujian pemulihan dan keluaran AI tetap memerlukan bukti tersendiri."
        ]
    ],
    "bukti_intro": "Bukti berikut menunjukkan dashboard yang tersedia dan struktur data saat pemeriksaan. Angka inventaris adalah snapshot 2 Oktober, sedangkan capaian pekerjaan dibatasi sampai 30 September.",
    "bukti": [
        [
            "img",
            "01-dashboard.jpg",
            "Gambar 1. Dashboard Statistik Pariwisata dan Ekonomi Kreatif dengan navigasi indikator."
        ],
        [
            "lake",
            "Inventaris struktur data saat verifikasi:"
        ],
        [
            "img",
            "06-katalog.jpg",
            "Gambar 2. Katalog data primer dan sekunder beserta penanda lapisan data."
        ]
    ],
    "rtl_peran": "Pada Oktober, Project Manager perlu menetapkan urutan pemenuhan data yang masih kosong bersama Dinas, mengoordinasikan pengujian angka antartampilan, dan memantau penyelesaian temuan mutu. Prioritas berikutnya adalah kelengkapan data dan ketertelusuran, bukan menambah indikator tanpa sumber.",
    "penutup": "Pekerjaan September memperluas pemanfaatan platform untuk statistik pariwisata dan ekonomi kreatif. Dashboard dan katalog telah diperiksa pada sistem hidup; kekosongan data dan keterbatasan bukti dicatat untuk tindak lanjut bulan berikutnya.",
    "rtl_bukti": []
},
{
    "filename": "02-Database-Administrator",
    "jabatan": "Tenaga Ahli Database Administrator (Senior)",
    "kualifikasi": "S1 Teknik Informatika/Ilmu Komputer/Sistem Informasi, pengalaman minimal 7 tahun",
    "uraian_tugas": [
        "Merancang struktur basis data statistik pariwisata pada platform Database (logical & physical data model);",
        "Menyusun skema tabel, indeks, constraint, view, dan partisi sesuai kebutuhan analitik;",
        "Melaksanakan loading data hasil pengumpulan, pembersihan, dan validasi ke dalam Database;",
        "Melakukan optimasi kinerja basis data (query tuning, indexing, statistik, partitioning);",
        "Menyusun tata kelola akses dan keamanan basis data (role, privilege, audit trail);",
        "Menyusun pedoman backup, recovery, dan pemeliharaan basis data;",
        "Mendokumentasikan basis data (ERD, kamus tabel, panduan operasional)."
    ],
    "pendahuluan_peran": "Database Administrator memusatkan pekerjaan bulan keempat pada keberlanjutan penyimpanan data yang melayani aplikasi v2, penataan akses, serta dokumentasi operasi. Struktur lake house yang telah dibangun menjadi landasan bagi penambahan sumber dan tampilan statistik September.",
    "pelaksanaan_intro": "Realisasi peran ini didukung dokumentasi pengoperasian dan pelepasan aplikasi yang diperbarui pada September, serta pemeriksaan inventaris dan isi mart. Keberhasilan pencadangan atau pemulihan tidak dinyatakan tanpa uji yang dapat diperiksa.",
    "realisasi": [
        [
            "Struktur basis data statistik",
            "Menjaga pemisahan data mentah, tampilan data bersih, dan mart penyaji sebagai jalur pembacaan aplikasi. Inventaris saat verifikasi membedakan tabel data, metadata, dan tabel persiapan."
        ],
        [
            "Skema dan pemuatan data",
            "Menelusuri tersedianya tampilan data baru untuk statistik wisatawan nusantara dan tenaga kerja ekonomi kreatif, serta memastikan periode datanya terbaca pada rekap."
        ],
        [
            "Kinerja dan akses basis data",
            "Mempertahankan pemisahan akun aplikasi hanya-baca dari akses operator sebagaimana panduan operasional. Verifikasi laporan menggunakan pembacaan data tanpa perubahan pada server."
        ],
        [
            "Pedoman operasi dan pemulihan",
            "Memperbarui rujukan handover dan runbook deploy yang tercatat pada September. Kesiapan uji pemulihan dicatat sebagai pekerjaan lanjutan, tanpa mengklaim keberhasilan pemulihan bulan ini."
        ],
        [
            "Dokumentasi basis data",
            "Menyajikan inventaris lapisan dan daftar mart disertai sumber kueri; mart historis wisatawan mancanegara diberi keterangan tahun 2014 agar tidak disalahgunakan sebagai data terbaru."
        ]
    ],
    "bukti_intro": "Tabel berikut berasal dari pembacaan inventaris dan isi mart pada 2 Oktober 2026. Jumlah objek saat pemeriksaan tidak diperlakukan sebagai inventaris per 30 September.",
    "bukti": [
        [
            "lake",
            "Struktur penyimpanan data:"
        ],
        [
            "mart",
            "Mart penyaji dan jumlah baris saat verifikasi:"
        ],
        [
            "img",
            "10-api-docs.jpg",
            "Gambar 1. Dokumentasi antarmuka data publik sebagai rujukan akses dataset."
        ]
    ],
    "rtl_peran": "DBA perlu memeriksa konsistensi mart historis dengan sumber dashboard terbaru, menyiapkan uji pemulihan yang terdokumentasi, dan menyelaraskan inventaris operasional dengan tabel persiapan sebelum perubahan berikutnya diterapkan.",
    "penutup": "Struktur penyimpanan dan mart telah ditelusuri tanpa mengubah keadaan produksi. Penataan dokumentasi September mendukung operasi bersama; kemampuan pemulihan tetap menjadi agenda pengujian yang harus dibuktikan.",
    "rtl_bukti": []
},
{
    "filename": "03-MDM-Specialist",
    "jabatan": "Tenaga Ahli Master Data Management (MDM) Specialist (Senior)",
    "kualifikasi": "S1 Sistem Informasi/Teknik Informatika/Ilmu Komputer/Statistik, pengalaman minimal 7 tahun",
    "uraian_tugas": [
        "Mengidentifikasi entitas master data pariwisata Jakarta dan menyusun struktur data master;",
        "Menyusun kebijakan dan prosedur MDM (data ownership, data stewardship, alur perubahan, persetujuan, dan audit);",
        "Melaksanakan pembentukan master data melalui konsolidasi, deduplikasi, verifikasi, dan validasi data dari berbagai sumber;",
        "Menyusun mekanisme pemeliharaan dan pemutakhiran master data secara berkala;",
        "Memastikan integrasi master data ke dalam database dan ketersediaannya bagi Dashboard melalui metadata."
    ],
    "pendahuluan_peran": "MDM Specialist memfokuskan bulan keempat pada konsistensi definisi indikator, periode data, dan keterlacakan sumber. Pengayaan dataset sekunder dan dashboard statistik menuntut pemisahan yang jelas antara inventori venue, jumlah kunjungan, serta ukuran tenaga kerja.",
    "pelaksanaan_intro": "Pelaksanaan mencakup penataan metadata yang terlihat pada katalog dan penelaahan hasil gerbang mutu yang memiliki riwayat September. Pemeriksaan membedakan jumlah temuan konversi dari jumlah baris sumber yang unik.",
    "realisasi": [
        [
            "Identifikasi entitas dan struktur master",
            "Mengelompokkan sumber wisatawan, usaha/venue pariwisata, dan tenaga kerja ekonomi kreatif agar ukuran yang berbeda tidak digabung sebagai satu indikator."
        ],
        [
            "Kebijakan dan kepemilikan data",
            "Menempatkan sumber, tahun/periode, dan batas ketersediaan data sebagai keterangan yang harus mengikuti angka pada dashboard. Data nasional per subsektor tidak diberi label sebagai rincian DKI."
        ],
        [
            "Konsolidasi dan validasi data",
            "Menelaah hasil gerbang mutu tanggal 30 September serta contoh alasan karantina. Temuan yang tersimpan diperlakukan sebagai bahan perbaikan, bukan bukti otomatis seluruh baris telah disisihkan."
        ],
        [
            "Pemeliharaan master data",
            "Memantau catatan jumlah baris selama September untuk menelusuri keberadaan pembaruan. Jumlah catatan tidak digunakan untuk menyimpulkan seluruh jadwal bulanan berhasil."
        ],
        [
            "Metadata untuk dashboard",
            "Mendukung kesesuaian judul, sumber resmi, cakupan wilayah, dan periode pada katalog serta halaman statistik tenaga kerja."
        ]
    ],
    "bukti_intro": "Riwayat gerbang mutu dan jumlah baris memiliki batas periode September. Karantina yang dapat dibaca merupakan snapshot 1 Oktober; keterbatasan itu dinyatakan pada catatan sumber.",
    "bukti": [
        [
            "qual",
            "Hasil gerbang mutu dan riwayat pembaruan:"
        ],
        [
            "karantina",
            "Temuan karantina yang tersedia saat pemeriksaan:"
        ],
        [
            "img",
            "06-katalog.jpg",
            "Gambar 1. Katalog dengan penanda lapisan dan sumber dataset."
        ],
        [
            "img",
            "05-tenaga-kerja.jpg",
            "Gambar 2. Dashboard tenaga kerja dengan pembedaan cakupan DKI dan nasional."
        ]
    ],
    "rtl_peran": "Pada Oktober, MDM Specialist perlu menyepakati definisi dan pemilik indikator yang masih kosong, menindaklanjuti kolom gagal konversi, serta menyiapkan penyimpanan riwayat karantina agar kondisi akhir bulan dapat ditelusuri.",
    "penutup": "Konsistensi definisi dan periode menjadi pengendalian utama dalam pengembangan September. Riwayat mutu mendukung penelaahan, sedangkan keterbatasan snapshot karantina perlu ditutup pada tahap berikutnya.",
    "rtl_bukti": []
},
{
    "filename": "04-BI-Developer-1",
    "jabatan": "Tenaga Ahli BI Developer 1 (Junior)",
    "kualifikasi": "S1 Sistem Informasi/Teknik Informatika/Ilmu Komputer/Statistik, pengalaman minimal 5 tahun",
    "uraian_tugas": [
        "Mengkonfigurasi koneksi Dashboard ke database melalui semantic layer/metadata MDM;",
        "Membangun dashboard interaktif (overview pariwisata, kunjungan wisatawan, akomodasi, destinasi, kontribusi ekonomi, dan dashboard tematik lainnya);",
        "Menyusun visualisasi yang konsisten dengan standar metadata dan definisi data;",
        "Mengatur akses pengguna, jadwal refresh data, dan pemeliharaan dashboard;",
        "Menyusun panduan pemanfaatan dashboard bagi pengguna di lingkungan Dinas."
    ],
    "pendahuluan_peran": "BI Developer mengembangkan tampilan statistik pada bulan keempat agar pengguna dapat membaca indikator pariwisata dan ekonomi kreatif dalam alur yang konsisten. Pengembangan September mencakup halaman ringkasan, rincian tahunan dan bulanan, serta komposisi indikator.",
    "pelaksanaan_intro": "Riwayat perubahan September menunjukkan pengembangan menu statistik, penyatuan tampilan wisman, pengayaan wisnus per kota tujuan, dan penataan visualisasi PDRB serta tenaga kerja. Halaman tersebut diperiksa kembali pada aplikasi publik.",
    "realisasi": [
        [
            "Koneksi dashboard ke sumber data",
            "Memakai data lake house untuk wisatawan mancanegara, perjalanan wisatawan nusantara menurut kota tujuan, dan statistik tenaga kerja; sumber dan cakupan ditampilkan pada halaman."
        ],
        [
            "Dashboard interaktif",
            "Menyusun halaman ringkasan statistik serta rincian wisman, wisnus, PDRB ekonomi kreatif, dan tenaga kerja dalam navigasi yang saling terhubung."
        ],
        [
            "Visualisasi dan definisi indikator",
            "Menyatukan bagian tahunan dan bulanan wisman dalam satu halaman gulir, menampilkan persentase komposisi, dan membatasi rincian tenaga kerja pada data resmi yang tersedia."
        ],
        [
            "Pemeliharaan antarmuka",
            "Memperbaiki navigasi global agar tidak meluber pada perangkat bergerak dan menghindari pemilihan koordinat sebagai ukuran grafik otomatis, sesuai perubahan yang tercatat September."
        ],
        [
            "Panduan pemanfaatan",
            "Menyertakan keterangan sumber, periode parsial, dan status menunggu data pada visualisasi, sehingga pengguna dapat menilai keterbandingan angka."
        ]
    ],
    "bukti_intro": "Gambar berikut diambil pada 2 Oktober dari halaman statistik yang pengembangannya tercatat pada September. Angka di dalamnya memiliki tahun data masing-masing, bukan realisasi September 2026.",
    "bukti": [
        [
            "img",
            "01-dashboard.jpg",
            "Gambar 1. Ringkasan dashboard statistik pariwisata dan ekonomi kreatif."
        ],
        [
            "img",
            "02-wisman.jpg",
            "Gambar 2. Wisman: capaian tahunan, data tahun berjalan, dan target."
        ],
        [
            "img",
            "03-wisnus.jpg",
            "Gambar 3. Perjalanan wisatawan nusantara menurut kota tujuan."
        ],
        [
            "img",
            "04-pdrb-ekraf.jpg",
            "Gambar 4. PDRB ekonomi kreatif dengan keterangan periode dan target."
        ],
        [
            "img",
            "05-tenaga-kerja.jpg",
            "Gambar 5. Tenaga kerja ekonomi kreatif dari sumber resmi."
        ]
    ],
    "rtl_peran": "BI Developer perlu menguji ulang keterbandingan periode dan penyaring, melengkapi rincian wilayah saat data diterima, serta memeriksa konsistensi visualisasi di desktop dan perangkat bergerak. Grafik tidak dilengkapi dengan angka perkiraan untuk mengisi data kosong.",
    "penutup": "Dashboard September telah memperluas cakupan informasi sekaligus memperjelas batas data yang tersedia. Tampilan yang diperiksa memberi jalur baca bagi pengguna dari ringkasan menuju sumber dan rincian indikator.",
    "rtl_bukti": []
},
{
    "filename": "05-Data-Engineer-2",
    "jabatan": "Tenaga Ahli Data Engineer 2 (Junior)",
    "kualifikasi": "S1 Teknik Informatika/Ilmu Komputer/Sistem Informasi, pengalaman minimal 6 tahun",
    "uraian_tugas": [
        "Merancang dan melaksanakan pipeline pengumpulan dan integrasi data (ETL/ELT) dari berbagai sumber ke dalam Database;",
        "Membersihkan, mentransformasi, dan menstandarisasi raw data agar siap digunakan untuk analisis dan visualisasi;",
        "Memastikan integrasi data multi-sumber dilakukan dengan menjaga kualitas dan keterlacakan data;",
        "Menerapkan standar metadata dan master data dalam proses integrasi data;",
        "Menjamin keamanan data dan kepatuhan terhadap tata kelola data pada proses integrasi."
    ],
    "pendahuluan_peran": "Data Engineer menitikberatkan bulan keempat pada penambahan sumber statistik resmi dan data sekunder ke jalur pengolahan yang sudah tersedia. Pekerjaan September mencakup data perjalanan wisnus, tenaga kerja ekonomi kreatif, venue tematik, serta penyiapan pipeline jumlah kamar hotel.",
    "pelaksanaan_intro": "Pelaksanaan ditelusuri dari perubahan skrip dan dataset pada September serta catatan pemantauan jumlah baris. Penyiapan pipeline jumlah kamar hotel tanggal 30 September dibedakan dari perluasan riset dan pendaftaran katalog yang tercatat pada Oktober.",
    "realisasi": [
        [
            "Pipeline pengumpulan dan integrasi",
            "Menambahkan data perjalanan wisnus menurut kota tujuan dan statistik tenaga kerja ekonomi kreatif dari sumber resmi ke jalur data. Menyiapkan pipeline jumlah kamar hotel pada akhir September."
        ],
        [
            "Pembersihan dan standardisasi",
            "Menyusun konversi data perjalanan wisata menjadi tabel berperiode, serta menjaga identitas wilayah dan jenis kelamin pada data tenaga kerja agar rekap tidak menghitung ganda."
        ],
        [
            "Integrasi multi-sumber",
            "Memperluas inventori sekunder wellness, wisata air, asosiasi/organisasi, desa wisata, dan hotel transit berdasarkan pekerjaan September; mempertahankan keterangan sumber pada dataset."
        ],
        [
            "Metadata dan ketertelusuran",
            "Menautkan dataset ke katalog dan mempertahankan catatan jumlah baris yang dapat ditelusuri selama September. Kueri laporan memakai batas waktu WIB yang dinyatakan eksplisit."
        ],
        [
            "Keamanan dan tata kelola",
            "Menggunakan jalur pembacaan untuk verifikasi laporan tanpa memicu pemuatan ulang atau perubahan server. Kebutuhan pemeriksaan mutu sumber baru dicatat sebagai tindak lanjut."
        ]
    ],
    "bukti_intro": "Riwayat jumlah baris dan hasil gerbang mutu menyediakan bukti periode September. Tangkapan layar menunjukkan dataset sekunder yang dapat dibuka saat verifikasi.",
    "bukti": [
        [
            "qual",
            "Catatan pengolahan data selama September:"
        ],
        [
            "lake",
            "Keluaran struktur penyimpanan saat pemeriksaan:"
        ],
        [
            "img",
            "07-wellness.jpg",
            "Gambar 1. Dataset venue wellness dengan keterangan asal dan metode pengumpulan."
        ]
    ],
    "rtl_peran": "Data Engineer perlu memeriksa pembaruan sumber baru secara terjadwal, melengkapi riwayat karantina, dan menguji pemuatan bertahap. Penambahan data kamar hotel diteruskan dengan pencatatan status sumber serta waktu pengambilan per hotel.",
    "penutup": "Pekerjaan September memperluas sumber yang dapat dimanfaatkan dashboard. Bukti riwayat pengolahan tersedia, sementara kelengkapan jadwal, pengujian pemulihan, dan riwayat karantina tetap perlu diperiksa secara khusus.",
    "rtl_bukti": []
},
{
    "filename": "06-Data-Analyst",
    "jabatan": "Tenaga Ahli Data Analyst (Intermediate)",
    "kualifikasi": "S1 Sistem Informasi/Teknik Informatika/Ilmu Komputer/Statistik, pengalaman minimal 6 tahun",
    "uraian_tugas": [
        "Menyiapkan dataset untuk pelaporan dan visualisasi sesuai kebutuhan Dinas;",
        "Menyusun laporan rutin (bulanan, triwulanan, tahunan) dan laporan ad-hoc yang relevan;",
        "Memastikan kualitas dan keakuratan data yang disajikan pada antarmuka pelaporan dan dashboard;",
        "Berkolaborasi dengan pengguna data untuk memvalidasi kebutuhan informasi."
    ],
    "pendahuluan_peran": "Data Analyst memusatkan bulan keempat pada penyiapan rekap indikator yang dapat dibandingkan secara benar. Fokusnya adalah kunjungan wisatawan mancanegara, perjalanan wisatawan nusantara, dan tenaga kerja ekonomi kreatif dengan tahun data yang dinyatakan eksplisit.",
    "pelaksanaan_intro": "Rekap disusun dari kueri ke tampilan data bersih yang juga mendasari dashboard. Angka tahunan dipisahkan dari periode tahun berjalan; data wisman dalam mart lama yang berperiode 2014 tidak digunakan sebagai pengganti data terbaru.",
    "realisasi": [
        [
            "Dataset pelaporan dan visualisasi",
            "Menyiapkan pembacaan wisatawan nusantara menurut kota tujuan serta tenaga kerja ekonomi kreatif DKI dari sumber resmi yang ditambahkan pada September."
        ],
        [
            "Laporan indikator rutin",
            "Menyajikan rekap wisman tahunan dan tahun berjalan, perjalanan wisnus tahunan dan periode Januari–Mei, serta tenaga kerja DKI tahun terakhir yang tersedia."
        ],
        [
            "Ketepatan dan kualitas angka",
            "Memastikan total tenaga kerja memakai kategori Total tanpa menjumlahkannya kembali dengan rincian laki-laki dan perempuan. Jumlah perjalanan tidak diberi label orang unik."
        ],
        [
            "Validasi kebutuhan informasi",
            "Mencatat kebutuhan lama menginap non-bintang, rincian PDRB per wilayah, dan tenaga kerja per subsektor DKI yang masih belum tersedia pada halaman."
        ],
        [
            "Dokumentasi metodologi",
            "Menjaga rujukan metodologis mengenai satuan, periode, dan keterbandingan indikator agar ringkasan tidak mengubah makna sumber."
        ]
    ],
    "bukti_intro": "Tabel rekap memuat periode data dan kueri sumber yang terverifikasi pada 2 Oktober. Bukti visual menunjukkan cara informasi itu disajikan pada dashboard.",
    "bukti": [
        [
            "wis",
            "Rekap indikator terpilih dengan periode eksplisit:"
        ],
        [
            "img",
            "02-wisman.jpg",
            "Gambar 1. Realisasi wisman dan keterangan cakupan tahun berjalan."
        ],
        [
            "img",
            "03-wisnus.jpg",
            "Gambar 2. Rekap perjalanan wisata menurut kota tujuan."
        ],
        [
            "img",
            "05-tenaga-kerja.jpg",
            "Gambar 3. Statistik tenaga kerja DKI dan keterangan batas rincian yang tersedia."
        ]
    ],
    "rtl_peran": "Data Analyst perlu menyelaraskan seri data yang berbeda periode, meminta rincian wilayah dan lama menginap yang belum tersedia, serta memeriksa kembali pembaruan angka sebelum menjadi bahan kebijakan. Perbandingan terhadap target tahunan harus memberi keterangan apabila realisasi masih parsial.",
    "penutup": "Rekap September mengutamakan keterlacakan dan keterbandingan, sehingga pembaca mengetahui ukuran serta tahun data setiap angka. Kelengkapan sumber berikutnya menentukan perluasan analisis yang dapat dipertanggungjawabkan.",
    "rtl_bukti": []
},
{
    "filename": "07-Business-Analyst-1",
    "jabatan": "Tenaga Ahli Business Analyst 1 (Intermediate)",
    "kualifikasi": "S1 Sistem Informasi/Teknik Informatika/Ilmu Komputer/Manajemen/Statistik, pengalaman minimal 6 tahun",
    "uraian_tugas": [
        "Melakukan pemetaan kondisi tata kelola data pariwisata saat ini (as-is) melalui wawancara, observasi, dan analisis dokumen;",
        "Merumuskan kondisi tata kelola data yang dituju (to-be) bersama pemangku kepentingan;",
        "Menyusun gap analysis as-is vs to-be beserta rekomendasi pemenuhan;",
        "Menerjemahkan kebutuhan bisnis menjadi spesifikasi tata kelola data (standar data, metadata, kebijakan MDM, kebijakan akses);",
        "Mendokumentasikan requirement, alur proses, dan use case pengelolaan data pariwisata;",
        "Berkoordinasi dengan pemangku kepentingan internal dan eksternal Dinas;",
        "Mendukung penyusunan roadmap dan rencana implementasi standarisasi data."
    ],
    "pendahuluan_peran": "Business Analyst memusatkan bulan keempat pada penerjemahan kebutuhan informasi menjadi susunan dashboard dan daftar kesenjangan data. Pengembangan September mempertemukan statistik pariwisata, ekonomi kreatif, inventori usaha, dan akses data dalam satu platform.",
    "pelaksanaan_intro": "Pelaksanaan ditelusuri melalui struktur menu statistik, keterangan kebutuhan data pada halaman, serta pembaruan dokumentasi handover dan runbook. Tampilan asisten AI dicatat sebagai antarmuka yang tersedia; ketepatan jawabannya memerlukan pengujian tersendiri.",
    "realisasi": [
        [
            "Pemetaan kondisi tata kelola",
            "Memetakan sumber yang tersedia pada katalog dan dashboard, termasuk data primer, data sekunder, serta sumber resmi statistik ekonomi kreatif."
        ],
        [
            "Perumusan kondisi yang dituju",
            "Mengarahkan kebutuhan pembacaan dari ringkasan pimpinan menuju rincian indikator, dengan sumber, periode, dan satuan yang tetap dapat diketahui."
        ],
        [
            "Analisis kesenjangan",
            "Mencatat rincian PDRB per wilayah, lama menginap non-bintang, serta tenaga kerja per subsektor DKI sebagai kebutuhan yang belum terisi, tanpa menggantinya dengan data nasional."
        ],
        [
            "Spesifikasi kebutuhan informasi",
            "Membedakan indikator kunjungan, perjalanan, inventori usaha/venue, dan tenaga kerja tersertifikasi agar kebutuhan bisnis tidak menghasilkan rekap yang keliru."
        ],
        [
            "Dokumentasi dan koordinasi",
            "Mendukung dokumentasi API, handover, serta runbook operasional yang diperbarui September sebagai dasar pemanfaatan dan pemeliharaan oleh tim."
        ],
        [
            "Roadmap implementasi",
            "Menempatkan pemenuhan data kosong dan pengujian jawaban AI sebagai agenda berikutnya sebelum memperluas pemanfaatan untuk analisis kebijakan."
        ]
    ],
    "bukti_intro": "Bukti menunjukkan struktur dashboard, antarmuka akses data, serta asisten data saat pemeriksaan. Ketersediaan antarmuka AI tidak diperlakukan sebagai bukti ketepatan keluaran analisis.",
    "bukti": [
        [
            "img",
            "01-dashboard.jpg",
            "Gambar 1. Struktur dashboard untuk kebutuhan statistik pariwisata dan ekonomi kreatif."
        ],
        [
            "img",
            "10-api-docs.jpg",
            "Gambar 2. Dokumentasi API untuk katalog dan pemanfaatan data."
        ],
        [
            "img",
            "09-ai.jpg",
            "Gambar 3. Antarmuka asisten data; jawaban analitis belum diuji dalam verifikasi laporan."
        ]
    ],
    "rtl_peran": "Business Analyst perlu menyepakati prioritas dan pemilik data yang belum tersedia bersama pengguna Dinas, merinci kebutuhan pengujian AI, serta menautkan indikator yang sudah valid ke pertanyaan kebijakan yang terukur.",
    "penutup": "Pemetaan kebutuhan September menghasilkan struktur informasi yang lebih jelas dan daftar kesenjangan yang dapat ditindaklanjuti. Pemanfaatan lanjutan perlu bertumpu pada kelengkapan sumber dan pengujian keluaran.",
    "rtl_bukti": []
},
]


def kop(doc, judul, subjudul):
    para(doc, "PEMERINTAH PROVINSI DAERAH KHUSUS IBUKOTA JAKARTA", size=12, bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    para(doc, "DINAS PARIWISATA DAN EKONOMI KREATIF", size=13, bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    para(doc, "Bidang Data, Informasi dan Pengembangan Destinasi", size=10, italic=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, color=GREY, space_after=2)
    hr = doc.add_paragraph(); hr.paragraph_format.space_after = Pt(10)
    pPr = hr._p.get_or_add_pPr(); pbdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom'); bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '12')
    bottom.set(qn('w:space'), '1'); bottom.set(qn('w:color'), '1a1a1a'); pbdr.append(bottom); pPr.append(pbdr)
    para(doc, judul, size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1)
    para(doc, subjudul, size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, color=GREY, space_after=10)


CAPAIAN_HEADER = ["Keluaran September", "Wujud dan bukti", "Status"]
CAPAIAN_ROWS = [
    [
        "Dashboard statistik pariwisata dan ekonomi kreatif",
        "Ringkasan serta halaman wisman, wisnus, PDRB, dan tenaga kerja; pengembangan tercatat 28–30 September",
        "Tersedia pada sistem publik"
    ],
    [
        "Integrasi statistik resmi",
        "Perjalanan wisnus menurut kota tujuan dan statistik tenaga kerja Kemenekraf; tahun data dan cakupan dinyatakan",
        "Data terverifikasi Q6–Q8"
    ],
    [
        "Pengayaan dataset sekunder",
        "Venue wellness, wisata air, asosiasi/organisasi, desa wisata, hotel transit; perubahan tercatat September",
        "Katalog diperiksa; inventaris snapshot"
    ],
    [
        "Pengayaan indikator kuliner",
        "Data mentah TripAdvisor ditambahkan pada halaman indikator kuliner tanggal 21 September",
        "Tampilan diperiksa"
    ],
    [
        "Asisten AI data",
        "Antarmuka asisten tersedia; pengembangan tercatat 17 September",
        "Antarmuka tersedia; jawaban belum diuji"
    ],
    [
        "Dokumentasi dan operasi bersama",
        "Handover dan runbook deploy diperbarui September; API dan buku metodologi tetap dapat diakses",
        "Dokumentasi tersedia"
    ],
    [
        "Penyiapan data jumlah kamar hotel",
        "Pipeline tercatat 30 September; riset lanjutan dan pendaftaran katalog bertanggal Oktober",
        "Penyiapan September; lanjutan di luar periode"
    ],
    [
        "Pemantauan mutu selama September",
        "Riwayat jumlah baris dan hasil gerbang mutu akhir bulan dapat ditelusuri melalui Q3–Q4",
        "Catatan tersedia; bukan bukti seluruh jadwal berhasil"
    ]
]
CAPAIAN_WIDTHS = [4.2, 8.0, 3.4]
KENDALA_HEADER = ["Kendala / risiko", "Dampak", "Tindakan dan mitigasi"]
KENDALA_ROWS = [
    [
        "Periode sumber tidak seragam",
        "Sebagian indikator berisi tahun 2024/2025 atau periode parsial 2026",
        "Cantumkan tahun, satuan, dan batas cakupan; hindari membandingkan periode parsial dengan target tahunan tanpa keterangan."
    ],
    [
        "Rincian yang belum tersedia",
        "Lama menginap non-bintang, PDRB per wilayah, dan tenaga kerja per subsektor DKI masih kosong",
        "Permintaan data kepada pemilik sumber; tampilkan status menunggu data, tanpa mengganti rincian DKI dengan data nasional."
    ],
    [
        "Mart historis dan tabel persiapan",
        "Mart wisman lama berperiode 2014; terdapat tabel _baru di lapisan penyaji",
        "Gunakan sumber dashboard terbaru untuk indikator; pisahkan mart operasional dari objek persiapan."
    ],
    [
        "Riwayat karantina terbatas",
        "Snapshot akhir September tidak tersedia pada tabel karantina yang dibaca",
        "Nyatakan tanggal snapshot yang tersedia dan siapkan penyimpanan riwayat sebagai tindak lanjut."
    ],
    [
        "Pengujian operasi dan AI",
        "Dokumentasi atau antarmuka belum membuktikan pemulihan cadangan dan ketepatan jawaban AI",
        "Jadwalkan uji tersendiri, catat hasil dan batasnya sebelum mengklaim kemampuan."
    ]
]
KENDALA_WIDTHS = [3.8, 5.0, 6.8]
RENCANA_HEADER = ["Rencana Oktober 2026", "Sasaran"]
RENCANA_ROWS = [
    [
        "Pemenuhan data indikator yang masih kosong",
        "Lama menginap non-bintang, rincian PDRB per wilayah, dan tenaga kerja per subsektor DKI diperoleh dari pemilik data."
    ],
    [
        "Penyelarasan sumber dashboard dan mart",
        "Periode, definisi, serta cakupan indikator konsisten; arsip historis diberi label yang jelas."
    ],
    [
        "Tindak lanjut temuan mutu dan riwayat karantina",
        "Temuan konversi diperiksa bersama pemilik sumber; snapshot per periode dapat ditelusuri."
    ],
    [
        "Uji pemulihan dan pembaruan terjadwal",
        "Kemampuan operasi dibuktikan dengan catatan pengujian, tidak disimpulkan dari dokumen konfigurasi."
    ],
    [
        "Uji jawaban asisten AI",
        "Pertanyaan terukur dibandingkan dengan hasil kueri dan sumber yang sama."
    ],
    [
        "Kelanjutan inventori jumlah kamar hotel",
        "Sumber, tanggal pengambilan, serta status verifikasi setiap hotel tercatat."
    ]
]
RENCANA_WIDTHS = [7.8, 7.8]
LAMPIRAN_HEADER = ["No.", "Posisi tenaga ahli", "Berkas September 2026"]
LAMPIRAN_ROWS = [[str(i), m["jabatan"], m["filename"] + ".docx / .pdf"] for i, m in enumerate(MEMBERS, 1)]
LAMPIRAN_WIDTHS = [1.0, 8.0, 6.6]
BUKTI_UTAMA = [
    ("01-dashboard.jpg", "Gambar 1. Ringkasan statistik pariwisata dan ekonomi kreatif."),
    ("02-wisman.jpg", "Gambar 2. Wisman dengan cakupan tahunan dan tahun berjalan."),
    ("03-wisnus.jpg", "Gambar 3. Perjalanan wisatawan nusantara menurut kota tujuan."),
    ("04-pdrb-ekraf.jpg", "Gambar 4. PDRB ekonomi kreatif dan pembanding target."),
    ("05-tenaga-kerja.jpg", "Gambar 5. Tenaga kerja ekonomi kreatif dari sumber resmi."),
    ("06-katalog.jpg", "Gambar 6. Katalog data primer dan sekunder dengan penanda lapisan."),
    ("07-wellness.jpg", "Gambar 7. Dataset sekunder venue wellness."),
    ("08-kuliner.jpg", "Gambar 8. Data pendukung indikator kuliner, termasuk TripAdvisor."),
    ("09-ai.jpg", "Gambar 9. Antarmuka asisten AI; jawaban analitis belum diuji."),
    ("10-api-docs.jpg", "Gambar 10. Dokumentasi antarmuka data publik."),
    ("11-buku-home.jpg", "Gambar 11. Buku Statistika Pariwisata Perkotaan sebagai rujukan metodologi."),
]


def build_utama():
    """Laporan utama (payung) kegiatan bulan ke-4 — merangkum tujuh laporan tenaga ahli."""
    doc = Document()
    st = doc.styles['Normal']; st.font.name = "Arial"; st.font.size = Pt(11)
    for section in doc.sections:
        section.top_margin = Cm(2.2); section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.5); section.right_margin = Cm(2.2)

    kop(doc,
        "LAPORAN BULANAN KEGIATAN",
        "Penyediaan dan Pengelolaan Data Statistik Pariwisata Jakarta — Tahun Anggaran 2026")

    t = doc.add_table(rows=0, cols=2); set_table_borders(t, color="bfbfbf")
    ident = [
        ("Periode Laporan", "Bulan ke-4 (September 2026)"),
        ("Kegiatan", "Analisis Data dan Pembangunan Dashboard Pariwisata"),
        ("Sub Kegiatan", "Perencanaan Daya Tarik Wisata Provinsi"),
        ("Lokasi", "Provinsi DKI Jakarta"),
        ("Jumlah Tenaga Ahli", "7 (tujuh) orang — laporan per orang terlampir"),
        ("Tanggal Verifikasi Angka", "2 Oktober 2026; riwayat mutu dibatasi sampai 30 September"),
    ]
    for k, v in ident:
        row = t.add_row().cells
        cell_text(row[0], k, bold=True, size=10, align=WD_ALIGN_PARAGRAPH.LEFT); set_cell_bg(row[0], "eeeeee")
        cell_text(row[1], v, size=10, align=WD_ALIGN_PARAGRAPH.LEFT)
        row[0].width = Cm(4.8); row[1].width = Cm(10.8)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


    heading(doc, "I.  RINGKASAN EKSEKUTIF")
    para(doc, PENDAHULUAN_UMUM)
    para(doc, "Inti capaian September adalah perluasan pemanfaatan data menjadi dashboard statistik yang menghubungkan wisatawan, ekonomi kreatif, dan sumber pendukung. Keterangan periode dan kekosongan data menjadi bagian dari keluaran, sehingga angka tidak ditafsirkan melebihi cakupannya.")

    heading(doc, "II.  CAPAIAN BULAN INI")
    para(doc, "Capaian berikut ditelusuri dari perubahan repositori selama September. Ketersediaan tampilan diperiksa pada 2 Oktober; perubahan Oktober dinyatakan terpisah.")
    data_table(doc, CAPAIAN_HEADER, CAPAIAN_ROWS, CAPAIAN_WIDTHS)
    para(doc, "Sumber pekerjaan: riwayat Git 1–30 September 2026 WIB; dokumentasi CLAUDE.md, docs/HANDOVER.md dan docs/DEPLOY-RUNBOOK.md. Bukti layar merujuk URL pada setiap caption.", size=8.5, italic=True, color=GREY, align=WD_ALIGN_PARAGRAPH.LEFT)

    heading(doc, "III.  STRUKTUR DATA YANG DIBANGUN")
    para(doc, "Lake house tiga lapis memisahkan data mentah, data bersih, dan mart penyaji. Inventaris berikut merupakan keadaan saat verifikasi 2 Oktober; jumlah objek pada akhir September tidak direkonstruksi dari snapshot ini.")
    data_table(doc, LAKE_HEADER, LAKE_ROWS, LAKE_WIDTHS)
    para(doc, TABLES["lake"][3], size=8.5, italic=True, color=GREY, align=WD_ALIGN_PARAGRAPH.LEFT)
    para(doc, "Mart penyaji yang tersedia:", bold=True, size=10.5, align=WD_ALIGN_PARAGRAPH.LEFT)
    data_table(doc, MART_HEADER, MART_ROWS, MART_WIDTHS)
    para(doc, TABLES["mart"][3], size=8.5, italic=True, color=GREY, align=WD_ALIGN_PARAGRAPH.LEFT)

    heading(doc, "IV.  PENJAMINAN MUTU DATA")
    para(doc, "Hasil gerbang mutu tanggal 30 September dan riwayat jumlah baris selama bulan tersebut dapat dibaca kembali. Hasil pemeriksaan gagal tetap dicatat; keberadaan catatan bukan bukti semua data sudah layak atau semua jadwal berhasil.")
    data_table(doc, QUAL_HEADER, QUAL_ROWS, QUAL_WIDTHS)
    para(doc, TABLES["qual"][3], size=8.5, italic=True, color=GREY, align=WD_ALIGN_PARAGRAPH.LEFT)
    data_table(doc, KARANTINA_HEADER, KARANTINA_ROWS, KARANTINA_WIDTHS)
    para(doc, TABLES["karantina"][3], size=8.5, italic=True, color=GREY, align=WD_ALIGN_PARAGRAPH.LEFT)

    heading(doc, "V.  BUKTI PELAKSANAAN")
    para(doc, "Tangkapan layar diambil pada 2 Oktober 2026 dari sistem yang berjalan. Gambar menunjukkan ketersediaan tampilan saat pemeriksaan, bukan snapshot kondisi 30 September. Tahun dan angka di dalam gambar mengikuti keterangan sumber pada halaman.")
    for filename, text in BUKTI_UTAMA:
        add_image(doc, filename, caption(filename, text))

    heading(doc, "VI.  REKAP INDIKATOR TERPILIH")
    para(doc, "Rekap berikut berasal dari tampilan data bersih yang mendasari dashboard. Data tahunan dan tahun berjalan dipisahkan; jumlah perjalanan bukan jumlah wisatawan unik.")
    data_table(doc, WIS_HEADER, WIS_ROWS, WIS_WIDTHS)
    para(doc, TABLES["wis"][3], size=8.5, italic=True, color=GREY, align=WD_ALIGN_PARAGRAPH.LEFT)
    para(doc, "Kueri lengkap dan hasil pemeriksaan tersimpan dalam verifikasi-bulan4.json. Q1 = inventory; Q2 = marts; Q3 = quality_september_last; Q4 = rowcount_history; Q5 = quarantine; Q6–Q8 = indicators menurut urutan sumber.", size=8.5, italic=True, color=GREY, align=WD_ALIGN_PARAGRAPH.LEFT)

    heading(doc, "VII.  KENDALA, RISIKO, DAN MITIGASI")
    data_table(doc, KENDALA_HEADER, KENDALA_ROWS, KENDALA_WIDTHS)

    heading(doc, "VIII.  RENCANA BULAN BERIKUTNYA")
    para(doc, RTL_UMUM)
    data_table(doc, RENCANA_HEADER, RENCANA_ROWS, RENCANA_WIDTHS)

    heading(doc, "IX.  PENUTUP")
    para(doc, "Pekerjaan bulan keempat memperluas dashboard statistik, sumber resmi, dataset sekunder, dan dokumentasi pemanfaatan platform. Angka dalam tabel diperiksa pada 2 Oktober 2026 dengan periode sumber yang dinyatakan. Riwayat September digunakan apabila tersedia; keterbatasan snapshot dan kemampuan yang belum diuji dicatat sebagai agenda Oktober.")

    heading(doc, "LAMPIRAN — LAPORAN PELAKSANAAN TUGAS TENAGA AHLI")
    para(doc, "Tujuh laporan tenaga ahli untuk bulan ke-4 September 2026 disertakan dalam direktori bulan-4, masing-masing dalam format DOCX dan PDF.")
    data_table(doc, LAMPIRAN_HEADER, LAMPIRAN_ROWS, LAMPIRAN_WIDTHS)
    # tanda tangan
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    sig = doc.add_table(rows=0, cols=2)
    r = sig.add_row().cells
    cell_text(r[0], "Mengetahui,\nPejabat Pembuat Komitmen", size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    cell_text(r[1], "Jakarta, ......................... 2026\nProject Manager Kegiatan",
              size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    for _ in range(4): sig.add_row()
    r = sig.add_row().cells
    cell_text(r[0], "(  Bima Agung  )", bold=True, size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    cell_text(r[1], "(  [Nama Project Manager]  )", bold=True, size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = sig.add_row().cells
    cell_text(r[0], "NIP. 197907162011011008", size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, color=GREY)
    cell_text(r[1], "Project Manager Analisis Data dan Pembangunan Dashboard Pariwisata",
              size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, color=GREY)

    outdir = os.path.join(BASE, "bulan-4")
    os.makedirs(outdir, exist_ok=True)
    out = os.path.join(outdir, "00-Laporan-Utama-Bulan-4.docx")
    prepare_layout(doc)
    doc.save(out)
    print("wrote", out)

if __name__ == "__main__":
    build_utama()
    for m in MEMBERS:
        build(m)
    print("done: 1 laporan utama +", len(MEMBERS), "laporan tenaga ahli")
