"""Ekspor (E §4.2 butir 6): CSV BOM/CRLF/kutip, XLSX sheet Data, nama berkas, 413, 404."""

import io

from openpyxl import load_workbook

from data_api.main import buat_app
from data_api.shared.konfigurasi import muat


def test_csv(klien):
    r = klien.get("/v1/dataset/wisman/ekspor")
    assert r.status_code == 200
    assert r.headers["content-type"] == "text/csv; charset=utf-8"
    assert r.headers["content-disposition"] == 'attachment; filename="wisman.csv"'
    assert r.headers["x-row-count"] == "5" and r.headers["cache-control"] == "no-store"
    teks = r.content.decode()
    assert teks.startswith("﻿periode_data,kebangsaan,jumlah,catatan\r\n")
    baris = teks.removeprefix("﻿").split("\r\n")
    assert len(baris) == 6 and not teks.endswith("\r\n")
    assert '2026-01-01,JEPANG,100,"a ""b"", c"' in baris
    assert "2026-02-01,JEPANG,200," in baris


def test_nama_terfilter(klien):
    for params in ({"q": "jepang"}, {"f": "kebangsaan:in:JEPANG"}, {"sort": "jumlah:asc"}):
        r = klien.get("/v1/dataset/wisman/ekspor", params=params)
        assert r.headers["content-disposition"] == 'attachment; filename="wisman-terfilter.csv"', params
    assert klien.get("/v1/dataset/wisman/ekspor", params={"q": "jepang"}).headers["x-row-count"] == "2"


def test_xlsx(klien):
    r = klien.get("/v1/dataset/wisman/ekspor", params={"format": "xlsx"})
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert r.headers["content-disposition"] == 'attachment; filename="wisman.xlsx"'
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Data"]
    rows = list(wb["Data"].iter_rows(values_only=True))
    assert rows[0] == ("periode_data", "kebangsaan", "jumlah", "catatan") and len(rows) == 6


def test_tidak_dikenal_404_dan_format_salah(klien):
    assert klien.get("/v1/dataset/tidak-ada/ekspor").status_code == 404
    assert klien.get("/v1/dataset/wisman/ekspor", params={"format": "pdf"}).status_code == 400


def test_terlalu_besar_413(ch, env):
    from fastapi.testclient import TestClient

    from tests.conftest import TOKEN

    app = buat_app(muat({**env, "EXPORT_MAX_ROWS": "4"}), ch)
    with TestClient(app, headers={"Authorization": f"Bearer {TOKEN}"}) as c:
        r = c.get("/v1/dataset/wisman/ekspor")
        assert r.status_code == 413 and r.json()["error"]["code"] == "TOO_LARGE"
        assert c.get("/v1/dataset/wisman/ekspor", params={"q": "jepang"}).status_code == 200


def test_ekspor_dibaca_per_halaman(klien, ch, monkeypatch):
    from data_api.features.dataset import service

    monkeypatch.setattr(service, "EKSPOR_PER_HALAMAN", 2)
    r = klien.get("/v1/dataset/wisman/ekspor")
    assert len(r.content.decode().split("\r\n")) == 6
    halaman = [p for s, p in ch.log if s.startswith("SELECT * FROM silver")]
    assert [(p["limit"], p["offset"]) for p in halaman] == [(2, 0), (2, 2), (2, 4)]


def test_sel_formula_dinetralkan():
    from data_api.features.dataset.service import _nilai_xlsx, _sel_csv

    for teks in ('=HYPERLINK("http://x")', "+1+1", "-2+3", "@SUM(A1)", "\tx", "\rx"):
        assert _sel_csv(teks).lstrip('"').startswith("'"), teks
        assert _nilai_xlsx(teks).startswith("'"), teks
    # angka (bertipe atau teks Int64 dari ClickHouse) dan teks biasa tidak diubah
    assert [_sel_csv(v) for v in (-20.34, "-5", "-1.5e3", "biasa", "a-b")] == ["-20.34", "-5", "-1.5e3", "biasa", "a-b"]
    assert _nilai_xlsx(-20.34) == -20.34 and _nilai_xlsx("-5") == "-5"


def test_ekspor_menetralkan_formula(klien, ch, monkeypatch):
    import copy
    import sys

    conftest = sys.modules[type(ch).__module__]  # modul conftest yang dipakai ChPalsu

    baris = copy.deepcopy(conftest.BARIS["t_wisman"])
    baris[1]["catatan"] = '=HYPERLINK("http://jahat","klik")'
    monkeypatch.setitem(conftest.BARIS, "t_wisman", baris)
    csv = klien.get("/v1/dataset/wisman/ekspor").content.decode()
    assert '"\'=HYPERLINK(""http://jahat"",""klik"")"' in csv
    r = klien.get("/v1/dataset/wisman/ekspor", params={"format": "xlsx"})
    nilai = [c for row in load_workbook(io.BytesIO(r.content))["Data"].iter_rows(values_only=True) for c in row]
    assert '\'=HYPERLINK("http://jahat","klik")' in nilai
